"""Update monthly coffee benchmarks (World Bank Pink Sheet, with FRED fallback).

Do not blend the two sources within one historical time series: their definitions differ.
"""
import csv
import io
import json
import re
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

OUTPUT = Path(__file__).with_name('data.json')
FRED = {'arabica': 'PCOFFOTMUSDM', 'robusta': 'PCOFFROBUSDM', 'exchange_rate': 'DEXKOUS'}
WORLD_BANK_URLS = [
    'https://thedocs.worldbank.org/en/doc/5d903e848db1d1b83e0ec8f744e55570-0350012021/related/CMO-Historical-Data-Monthly.xlsx',
    'https://thedocs.worldbank.org/en/doc/18675f1d1639c7a34d463f59263ba0a2-0050012025/related/CMO-Historical-Data-Monthly.xlsx',
]
POUNDS_PER_KG = 2.2046226218


def download(url):
    req = Request(url, headers={'User-Agent': 'Mozilla/5.0 (compatible; CoffeePriceIndex/2.0)', 'Accept': '*/*'})
    with urlopen(req, timeout=45) as res:
        return res.read()


def parse_fred_csv(text, series):
    result = {}
    for row in csv.DictReader(io.StringIO(text)):
        key = row.get('observation_date') or row.get('DATE')
        raw = row.get(series)
        if key and raw and raw != '.':
            try:
                result[key] = round(float(raw), 5)
            except ValueError:
                pass
    return result


def fetch_fred(series):
    return parse_fred_csv(download(f'https://fred.stlouisfed.org/graph/fredgraph.csv?id={series}').decode('utf-8-sig'), series)


def month_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-01')
    text = str(value or '').strip()
    match = re.fullmatch(r'(\d{4})[Mm](0[1-9]|1[0-2])', text)
    if match:
        return f'{match.group(1)}-{match.group(2)}-01'
    for pattern in ('%Y-%m-%d', '%Y-%m', '%Y%m', '%b-%y', '%b %Y', '%Y %b'):
        try:
            return datetime.strptime(text, pattern).strftime('%Y-%m-01')
        except ValueError:
            pass
    return None


def parse_pink_sheet(blob):
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        sheet = next((s for s in wb if 'monthly' in s.title.lower()), wb.active)
        iterator = sheet.iter_rows(values_only=True)
        first = [next(iterator, ()) for _ in range(9)]
        headers = None
        for index, row in enumerate(first):
            columns = {re.sub(r'[^a-z]', '', str(v).lower()): j for j, v in enumerate(row) if v is not None}
            arabica = next((j for name, j in columns.items() if name in ('coffeearabica', 'coffeearabicas')), None)
            robusta = next((j for name, j in columns.items() if name == 'coffeerobusta'), None)
            if arabica is not None and robusta is not None:
                headers = (index, arabica, robusta)
                break
        if headers is None:
            raise ValueError('World Bank 커피 열을 찾지 못했습니다 (Coffee, Arabica / Coffee, Robusta)')
        idx, a_col, r_col = headers
        result = []
        for row in first[idx + 1:] + list(iterator):
            if len(row) <= max(a_col, r_col) or not row:
                continue
            day = month_date(row[0])
            try:
                a, r = float(row[a_col]), float(row[r_col])
            except (ValueError, TypeError):
                continue
            if day and 0 < a < 100 and 0 < r < 100:
                # World Bank source is USD/kg; site uses US cents/lb.
                result.append({'date': day, 'arabica': round(a * 100 / POUNDS_PER_KG, 5),
                               'robusta': round(r * 100 / POUNDS_PER_KG, 5)})
        result = list({item['date']: item for item in result}.values())
        result.sort(key=lambda item: item['date'])
        if len(result) < 100:
            raise ValueError(f'World Bank 원본 자료 행 부족: {len(result)}')
        return result
    finally:
        wb.close()


def fetch_world_bank():
    failures = []
    for url in WORLD_BANK_URLS:
        try:
            rows = parse_pink_sheet(download(url))
            return rows, url
        except Exception as exc:
            failures.append(f'{url}: {exc}')
    raise RuntimeError('World Bank 다운로드 실패: ' + ' | '.join(failures))


def fetch_fred_coffee():
    a, r = fetch_fred(FRED['arabica']), fetch_fred(FRED['robusta'])
    dates = sorted(set(a) & set(r))
    if not dates:
        raise RuntimeError('FRED 월별 공통 데이터가 없습니다.')
    return [{'date': d, 'arabica': a[d], 'robusta': r[d]} for d in dates]


def main():
    old = json.loads(OUTPUT.read_text(encoding='utf-8')) if OUTPUT.exists() else {}
    candidates = []
    problems = []
    try:
        rows, url = fetch_world_bank()
        candidates.append((rows[-1]['date'], rows, 'World Bank Commodity Price Data (Pink Sheet)', url))
    except Exception as exc:
        problems.append(str(exc))
    try:
        rows = fetch_fred_coffee()
        candidates.append((rows[-1]['date'], rows, 'International Monetary Fund via FRED',
                           'https://fred.stlouisfed.org/series/PCOFFOTMUSDM'))
    except Exception as exc:
        problems.append(str(exc))

    if not candidates:
        raise RuntimeError('커피 가격 다운로드가 모두 실패했습니다. 기존 파일 유지. ' + ' | '.join(problems))
    candidates.sort(key=lambda c: (c[0], c[2].startswith('World Bank')), reverse=True)
    latest, rows, source, source_url = candidates[0]
    if latest < old.get('latestObservation', ''):
        raise RuntimeError(f'기존 자료보다 오래된 데이터입니다: {latest} < {old["latestObservation"]}')
    cutoff = f'{int(latest[:4]) - 10:04d}-{latest[5:7]}-01'
    rows = [r for r in rows if r['date'] >= cutoff]
    if len(rows) < 100:
        raise RuntimeError('10년치 자료가 충분하지 않아 기존 데이터를 보존했습니다.')

    rate, rate_date = old.get('exchangeRate'), old.get('exchangeRateDate')
    try:
        fx = fetch_fred(FRED['exchange_rate'])
        d = max(fx)
        if not rate_date or d >= rate_date:
            rate, rate_date = fx[d], d
    except Exception as exc:
        problems.append('환율: ' + str(exc))
    if not rate or not rate_date:
        raise RuntimeError('사용 가능한 환율 정보가 없습니다.')

    payload = {
        'updatedAt': datetime.now(timezone.utc).isoformat(timespec='seconds'),
        'latestObservation': latest, 'exchangeRate': rate, 'exchangeRateDate': rate_date,
        'unit': 'U.S. cents per pound', 'frequency': 'Monthly',
        'source': source, 'sourceUrl': source_url,
        'series': FRED if 'FRED' in source else {'arabica': 'Coffee, Arabica', 'robusta': 'Coffee, Robusta', 'exchange_rate': 'DEXKOUS'},
        'rows': rows,
    }
    # Avoid weekly no-op commits just because the script was executed again.
    comparable = lambda obj: {k: v for k, v in obj.items() if k != 'updatedAt'}
    if comparable(payload) == comparable(old):
        print(f'No new observations ({latest}); data.json unchanged.')
        return
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Updated {len(rows)} months through {latest} from {source}; exchange {rate_date}')
    for problem in problems:
        print('Warning:', problem)


if __name__ == '__main__':
    main()
